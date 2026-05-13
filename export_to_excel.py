"""
SAP HANA Data Exporter to Excel.
Connects to live SAP HANA database, pulls all data from tables defined in BSL mapping,
and exports everything to a formatted Excel workbook (wfm_schema_export.xlsx).

Dependencies:
- openpyxl
- sqlalchemy
- sqlalchemy-hana
- python-dotenv
"""

import os
import sys
import datetime
from decimal import Decimal

# Ensure imports resolve correctly from repo root
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

import dotenv
from sqlalchemy import text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Local imports
from src.infrastructure.hana_connection import get_engine
from src.engine.bsl_dictionary import BSL_MAPPING

# Load credentials
dotenv.load_dotenv()

# --- STYLING CONSTANTS ---
# Domain color scheme
COLORS = {
    'HR': {
        'title_bg': '1F4E79',     # Dark Blue
        'light_bg': 'D6E4F0',     # Light Blue
        'alt_row_bg': 'EBF3FB',   # Very Light Blue
    },
    'OPS': {
        'title_bg': '1C5E1C',     # Dark Green
        'light_bg': 'D6F0D6',     # Light Green
        'alt_row_bg': 'EBFAEB',   # Very Light Green
    },
    'UNKNOWN': {
        'title_bg': '3A3A3A',     # Dark Gray
        'light_bg': 'EFEFEF',     # Light Gray
        'alt_row_bg': 'F9F9F9',   # Off White
    },
    'HEADER': '2E75B6',          # Medium Blue
    'WARNING_BG': 'FFF2CC',      # Yellow
    'WARNING_FONT': '7F4F00',     # Brownish
    'ERROR_BG': 'FFD7D7',        # Red
    'ERROR_FONT': '8B0000',      # Dark Red
    'PK_BG': 'FFF2CC',           # Light Yellow
    'WHITE': 'FFFFFF',
}

THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

def get_domain_colors(domain):
    return COLORS.get(domain.upper(), COLORS['UNKNOWN'])

def apply_full_border(cell):
    cell.border = THIN_BORDER

def set_sheet_name(bsl_key):
    """Truncate and clean sheet names to meet Excel limits (31 chars)."""
    name = bsl_key.replace('::', '__').replace('views_', 'vw_')
    if len(name) > 31:
        name = name[:31]
    return name

def auto_size_columns(ws, min_width=12, max_width=60, sample_rows=100):
    """Auto-size columns based on content."""
    for col in ws.columns:
        max_length = 0
        column = get_column_letter(col[0].column) # Get the column name
        
        # Check header and first few rows
        for i, cell in enumerate(col):
            if i > sample_rows:
                break
            try:
                if cell.value:
                    length = len(str(cell.value))
                    if length > max_length:
                        max_length = length
            except:
                pass
        
        adjusted_width = max(min_width, min(max_length + 2, max_width))
        ws.column_dimensions[column].width = adjusted_width

def create_index_sheet(wb, tables_config):
    """Create the initial INDEX sheet structure."""
    ws = wb.active
    ws.title = "INDEX"
    
    # Title row
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    title = f"WFM Agent — Schema Export (Generated: {now})"
    ws.merge_cells('A1:I1')
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=14, color=COLORS['WHITE'])
    ws['A1'].fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Header row
    headers = ["#", "BSL Key", "Physical Name", "Domain", "Primary Keys", "Total Columns", "Total Rows Fetched", "Warnings", "Description"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color=COLORS['WHITE'])
        cell.fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
        cell.border = THIN_BORDER
        
    return ws

def export_table_to_sheet(wb, bsl_key, table_info, engine, index_ws, row_idx):
    """Fetch data for a single table and write it to a new sheet."""
    physical_name = table_info['physical_name']
    domain = table_info['domain']
    description = table_info['description']
    bsl_pks = table_info.get('primary_keys', [])
    warnings_list = table_info.get('warnings', [])
    
    sheet_name = set_sheet_name(bsl_key)
    # Ensure uniqueness
    base_name = sheet_name
    counter = 2
    while sheet_name in wb.sheetnames:
        suffix = f"_{counter}"
        sheet_name = base_name[:31-len(suffix)] + suffix
        counter += 1
        
    ws = wb.create_sheet(title=sheet_name)
    colors = get_domain_colors(domain)
    
    print(f"[{row_idx-2}/25] Fetching: {physical_name} ... ", end="", flush=True)
    
    data = []
    columns_metadata = []
    error_msg = None
    
    # Split schema and table for proper quoting: "SCHEMA"."TABLE"
    if "." in physical_name:
        parts = physical_name.split(".", 1)
        quoted_name = f'"{parts[0]}"."{parts[1]}"'
    else:
        quoted_name = f'"{physical_name}"'

    try:
        with engine.connect() as conn:
            result = conn.execute(text(f'SELECT * FROM {quoted_name}'))
            
            # Extract column metadata from cursor before fetching rows
            if hasattr(result.cursor, 'description') and result.cursor.description:
                for idx, col_desc in enumerate(result.cursor.description):
                    columns_metadata.append({
                        'name': col_desc[0],
                        'type': str(col_desc[1]) if col_desc[1] else "UNKNOWN",
                        'nullable': "Yes" if (len(col_desc) > 6 and col_desc[6]) else "Unknown",
                        'is_pk': "Yes" if col_desc[0] in bsl_pks else "No",
                        'index': idx
                    })
            
            data = result.fetchall()
                
    except Exception as e:
        error_msg = str(e)
    
    # Final console status report
    num_rows = len(data)
    num_cols = len(columns_metadata) if columns_metadata else 0
    
    if error_msg:
        if "without input arguments" in error_msg:
            print(f"SKIPPED (Parameterized View: Requires arguments)")
            error_msg = "This is a parameterized view and requires input arguments to fetch live data."
        else:
            print(f"ERROR: {error_msg}")
    else:
        print(f"{num_rows} rows fetched. Sheet written.")
    
    # --- Section A: Table Metadata block ---
    # Row 1: Physical Name | Domain | BSL Key
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(5, num_cols))
    cell_a1 = ws['A1']
    cell_a1.value = f"{physical_name} | {domain} | {bsl_key}"
    cell_a1.font = Font(name='Calibri', bold=True, size=12, color=COLORS['WHITE'])
    cell_a1.fill = PatternFill(start_color=colors['title_bg'], end_color=colors['title_bg'], fill_type='solid')
    cell_a1.alignment = Alignment(horizontal='left', vertical='center')
    apply_full_border(cell_a1)
    
    # Row 2: Description
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(5, num_cols))
    cell_a2 = ws['A2']
    cell_a2.value = description
    cell_a2.font = Font(name='Calibri', size=10)
    cell_a2.fill = PatternFill(start_color=colors['light_bg'], end_color=colors['light_bg'], fill_type='solid')
    cell_a2.alignment = Alignment(wrap_text=True, vertical='center')
    apply_full_border(cell_a2)
    
    # Row 3: Primary Keys
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(5, num_cols))
    cell_a3 = ws['A3']
    pk_str = ", ".join(bsl_pks) if bsl_pks else "None"
    cell_a3.value = f"Primary Keys: {pk_str}"
    cell_a3.font = Font(name='Calibri', size=10)
    apply_full_border(cell_a3)
    
    # Row 4: Warnings
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max(5, num_cols))
    cell_a4 = ws['A4']
    if warnings_list:
        cell_a4.value = f"Warnings: {' '.join(warnings_list)}"
        cell_a4.fill = PatternFill(start_color=COLORS['WARNING_BG'], end_color=COLORS['WARNING_BG'], fill_type='solid')
        cell_a4.font = Font(name='Calibri', size=10, color=COLORS['WARNING_FONT'])
    else:
        cell_a4.value = "Warnings: None"
        cell_a4.font = Font(name='Calibri', size=10)
    apply_full_border(cell_a4)
        
    # Row 5: Spacer
    ws.row_dimensions[5].height = 8
    
    # --- Section B: Column metadata table ---
    meta_headers = ["Column Name", "Data Type", "Nullable", "Is Primary Key", "Live Column Index"]
    for c_idx, h in enumerate(meta_headers, 1):
        cell = ws.cell(row=6, column=c_idx)
        cell.value = h
        cell.font = Font(name='Calibri', bold=True, size=10, color=COLORS['WHITE'])
        cell.fill = PatternFill(start_color=COLORS['HEADER'], end_color=COLORS['HEADER'], fill_type='solid')
        apply_full_border(cell)
        
    if error_msg and not columns_metadata:
        # If fetch failed, use BSL for column names if possible, but mark types as N/A
        bsl_cols = table_info.get('columns', {})
        for r_idx, (col_name, col_info) in enumerate(bsl_cols.items(), 7):
            row_data = [col_name, "N/A", "N/A", "Yes" if col_name in bsl_pks else "No", "N/A"]
            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.font = Font(name='Calibri', size=10)
                apply_full_border(cell)
                # PK highlighting
                if col_name in bsl_pks:
                    cell.fill = PatternFill(start_color=COLORS['PK_BG'], end_color=COLORS['PK_BG'], fill_type='solid')
                elif (r_idx - 7) % 2 == 1: # Alternating
                    cell.fill = PatternFill(start_color=colors['alt_row_bg'], end_color=colors['alt_row_bg'], fill_type='solid')
        last_meta_row = 6 + len(bsl_cols)
    else:
        for r_idx, col_m in enumerate(columns_metadata, 7):
            row_data = [col_m['name'], col_m['type'], col_m['nullable'], col_m['is_pk'], col_m['index']]
            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.font = Font(name='Calibri', size=10)
                apply_full_border(cell)
                # PK highlighting
                if col_m['is_pk'] == "Yes":
                    cell.fill = PatternFill(start_color=COLORS['PK_BG'], end_color=COLORS['PK_BG'], fill_type='solid')
                elif (r_idx - 7) % 2 == 1: # Alternating
                    cell.fill = PatternFill(start_color=colors['alt_row_bg'], end_color=colors['alt_row_bg'], fill_type='solid')
        last_meta_row = 6 + len(columns_metadata)

    # --- Section C: Spacer ---
    spacer_row = last_meta_row + 1
    ws.row_dimensions[spacer_row].height = 10
    
    # --- Section D: Live Data section ---
    data_header_start = spacer_row + 1
    
    # Section D Header
    ws.merge_cells(start_row=data_header_start, start_column=1, end_row=data_header_start, end_column=max(1, num_cols))
    cell_dh = ws.cell(row=data_header_start, column=1)
    cell_dh.value = f"Live Data — {physical_name} — {num_rows} rows"
    cell_dh.font = Font(name='Calibri', bold=True, size=12, color=COLORS['WHITE'])
    cell_dh.fill = PatternFill(start_color=colors['title_bg'], end_color=colors['title_bg'], fill_type='solid')
    apply_full_border(cell_dh)
    
    # Data Column Headers
    data_col_header_row = data_header_start + 1
    if columns_metadata:
        for c_idx, col_m in enumerate(columns_metadata, 1):
            cell = ws.cell(row=data_col_header_row, column=c_idx)
            cell.value = col_m['name']
            cell.font = Font(name='Calibri', bold=True, size=10, color=COLORS['WHITE'])
            cell.fill = PatternFill(start_color=COLORS['HEADER'], end_color=COLORS['HEADER'], fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            apply_full_border(cell)
    
    # Freeze panes
    ws.freeze_panes = ws.cell(row=data_col_header_row + 1, column=1)
    
    # Data Rows
    if error_msg:
        ws.merge_cells(start_row=data_col_header_row + 1, start_column=1, end_row=data_col_header_row + 3, end_column=max(5, num_cols))
        err_cell = ws.cell(row=data_col_header_row + 1, column=1)
        err_cell.value = f"FETCH ERROR: {error_msg}"
        err_cell.fill = PatternFill(start_color=COLORS['ERROR_BG'], end_color=COLORS['ERROR_BG'], fill_type='solid')
        err_cell.font = Font(name='Calibri', color=COLORS['ERROR_FONT'], bold=True, size=10)
        err_cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        apply_full_border(err_cell)
    else:
        for r_idx, row_vals in enumerate(data, data_col_header_row + 1):
            ws.row_dimensions[r_idx].height = 15
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = Font(name='Calibri', size=10)
                
                # Handle types
                if val is None:
                    cell.value = ""
                elif isinstance(val, (datetime.datetime, datetime.date)):
                    cell.value = val
                elif isinstance(val, Decimal):
                    cell.value = float(val)
                elif isinstance(val, (bytes, memoryview)):
                    # Convert binary data to hex string for Excel
                    cell.value = f"0x{val.hex()}" if hasattr(val, 'hex') else f"0x{bytes(val).hex()}"
                else:
                    cell.value = val
                
                apply_full_border(cell)
                # Alternating row color
                if (r_idx - (data_col_header_row + 1)) % 2 == 1:
                    cell.fill = PatternFill(start_color=colors['alt_row_bg'], end_color=colors['alt_row_bg'], fill_type='solid')

    # Update INDEX sheet row
    index_ws.cell(row=row_idx, column=1).value = row_idx - 2
    index_ws.cell(row=row_idx, column=2).value = bsl_key
    index_ws.cell(row=row_idx, column=3).value = physical_name
    index_ws.cell(row=row_idx, column=4).value = domain
    index_ws.cell(row=row_idx, column=5).value = pk_str
    index_ws.cell(row=row_idx, column=6).value = num_cols
    index_ws.cell(row=row_idx, column=7).value = num_rows
    index_ws.cell(row=row_idx, column=8).value = " | ".join(warnings_list) if warnings_list else ""
    index_ws.cell(row=row_idx, column=9).value = description
    
    # Style index row
    row_fill = PatternFill(start_color=colors['light_bg'], end_color=colors['light_bg'], fill_type='solid')
    for c in range(1, 10):
        cell = index_ws.cell(row=row_idx, column=c)
        cell.fill = row_fill
        apply_full_border(cell)

    # Auto-size columns for the table sheet
    auto_size_columns(ws)

def main():
    """Main execution entry point."""
    print("--- SAP HANA Schema Export Tool ---")
    
    # Initialize workbook
    wb = Workbook()
    index_ws = create_index_sheet(wb, BSL_MAPPING['tables'])
    
    try:
        engine = get_engine()
    except Exception as e:
        print(f"CRITICAL ERROR: Could not connect to SAP HANA. {e}")
        return

    # Iterate through tables
    tables = BSL_MAPPING['tables']
    current_row = 3
    for bsl_key, table_info in tables.items():
        export_table_to_sheet(wb, bsl_key, table_info, engine, index_ws, current_row)
        current_row += 1
        
    # Finalize INDEX sheet
    auto_size_columns(index_ws)
    
    # Save file
    output_file = "wfm_schema_export.xlsx"
    try:
        wb.save(output_file)
        print(f"\nDone. File written to: {output_file}")
    except Exception as e:
        print(f"\nERROR: Could not save workbook. {e}")

if __name__ == "__main__":
    main()